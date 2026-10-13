"""build_pptx.py: a synthetic spec with every element type becomes the expected native objects."""
import json
import os
import tempfile
import unittest

from PIL import Image, ImageDraw

from helpers import NS, run, shapes_by_name, validate_pptx, xfrm_of

K = 12700  # EMU per source pixel for a 960 px wide source (12192000 / 960)

SPEC = {
    'size': [960, 540], 'source': 'source.png', 'background': '#FFFFFF', 'lang': 'fr-FR', 'title': 'Synthetic',
    'elements': [
        {'id': 'title', 'type': 'text', 'lines': ['Bonjour & <bienvenue>'], 'x': 60, 'baseline': 80, 'size': 32,
         'font': 'Arial', 'bold': True, 'color': '#1F2A44', 'lsb': 2},
        {'id': 'body', 'type': 'text', 'lines': ['ligne 1', 'ligne 2'], 'x': 60, 'baseline': 140, 'size': 16,
         'pitch': 18, 'font': 'Arial', 'lsb': 1},
        {'id': 'para', 'type': 'text', 'lines': ['deux', 'lignes'], 'x': 300, 'baseline': 140, 'size': 16,
         'pitch': 24, 'font': 'Arial', 'align': 'right', 'w': 60, 'rsb': 1},
        {'id': 'mixed', 'type': 'text', 'lines': [[{'text': 'gras ', 'bold': True}, {'text': 'normal'}]], 'x': 60,
         'baseline': 180, 'size': 16, 'font': 'Arial', 'spacing': 1},
        {'id': 'list', 'type': 'bullets', 'items': ['un', ['deux', 'suite']], 'x': 60, 'text_x': 80,
         'baseline': 220, 'size': 16, 'pitch': 24, 'gap': 6, 'bullet': '●', 'bullet_color': '#F26B3A',
         'bullet_size': 125, 'font': 'Arial'},
        {'id': 'card', 'type': 'rect', 'box': [500, 60, 900, 260], 'fill': '#F2F2F2', 'radius': 12,
         'line': {'color': '#DDDDDD', 'width': 1}},
        {'id': 'pill', 'type': 'rect', 'box': [500, 300, 700, 340], 'fill': '#1F2A44', 'radius': 'pill',
         'text': {'lines': ['Statut'], 'size': 14, 'color': '#FFFFFF', 'font': 'Arial'}},
        {'id': 'dot', 'type': 'ellipse', 'box': [720, 300, 760, 340], 'fill': '#F26B3A', 'opacity': 0.5},
        {'id': 'step1', 'type': 'arrow', 'box': [60, 400, 260, 460], 'tip': 20, 'fill': '#1F2A44'},
        {'id': 'step2', 'type': 'arrow', 'box': [250, 400, 450, 460], 'tip': 20, 'notch': True, 'fill': '#3A4A6B'},
        {'id': 'tri', 'type': 'polygon', 'points': [[800, 400], [900, 400], [850, 480]], 'fill': '#00AA55'},
        {'id': 'rule', 'type': 'line', 'p1': [60, 380], 'p2': [900, 380], 'width': 2, 'color': '#CCCCCC',
         'dash': 'dash', 'head': 'end'},
        {'id': 'logo', 'type': 'image', 'box': [880, 20, 940, 50], 'transparent': 'auto', 'alt': 'logo'},
        {'id': 'ghost', 'type': 'rect', 'box': [0, 0, 10, 10], 'fill': '#000000', 'skip': True},
    ]}


class BuildTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        d = cls.tmp.name
        img = Image.new('RGB', (960, 540), 'white')
        ImageDraw.Draw(img).ellipse([890, 25, 930, 45], fill=(200, 30, 30))
        img.save(os.path.join(d, 'source.png'))
        with open(os.path.join(d, 'spec.json'), 'w', encoding='utf-8') as f:
            json.dump(SPEC, f, ensure_ascii=False)
        cls.out = os.path.join(d, 'out', 'synthetic.pptx')
        cls.log = run('build_pptx.py', os.path.join(d, 'spec.json'), '--out', cls.out)
        cls.slide = validate_pptx(cls.out)['ppt/slides/slide1.xml']
        cls.sh = shapes_by_name(cls.slide)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def geom(self, name):
        return self.sh[name].find('p:spPr/a:prstGeom', NS)

    def test_slide_is_standard_16_9(self):
        import zipfile
        with zipfile.ZipFile(self.out) as z:
            pres = z.read('ppt/presentation.xml').decode('utf-8')
        self.assertIn('<p:sldSz cx="12192000" cy="6858000"/>', pres)

    def test_single_line_uses_the_baseline_rule(self):
        # one line, size S = 32 px, default pitch 1.3 S: P = 41.6 px, top = baseline - 0.75 P, left = x - lsb
        P = 1.3 * 32
        x, y, cx, cy = xfrm_of(self.sh['title'])
        self.assertEqual((x, y, cy), (58 * K, round((80 - 0.75 * P) * K), round(P * K)))
        sp = self.sh['title'].find('.//a:lnSpc/a:spcPts', NS)
        self.assertEqual(sp.get('val'), '4160')  # 41.6 px = 41.6 pt on a 960 px source
        self.assertIn('Bonjour & <bienvenue>', [t.text for t in self.sh['title'].iter(f'{{{NS["a"]}}}t')])
        self.assertTrue(all(r.get('lang') == 'fr-FR' for r in self.sh['title'].iter(f'{{{NS["a"]}}}rPr')))

    def test_tight_block_is_split_per_line(self):
        g = self.sh['body-lines']
        self.assertEqual(g.tag, f'{{{NS["p"]}}}grpSp')
        y1 = xfrm_of(self.sh['body-l1'])[1]
        y2 = xfrm_of(self.sh['body-l2'])[1]
        self.assertEqual(y2 - y1, 18 * K)  # the measured pitch, not PowerPoint's
        self.assertIn('line pitch', self.log)

    def test_loose_block_is_one_box(self):
        x, y, cx, cy = xfrm_of(self.sh['para'])
        self.assertEqual((y, cy), (140 * K - 18 * K, 48 * K))  # top = 140 - 0.75 x 24, height = 2 x 24
        self.assertEqual(x + cx, 301 * K)  # right aligned: ink edge on x, box edge one right bearing further
        self.assertEqual(self.sh['para'].find('.//a:pPr', NS).get('algn'), 'r')

    def test_mixed_runs_and_spacing(self):
        runs = self.sh['mixed'].findall('.//a:r', NS)
        self.assertEqual([r.find('a:t', NS).text for r in runs], ['gras ', 'normal'])
        self.assertEqual([r.find('a:rPr', NS).get('b') for r in runs], ['1', None])
        self.assertEqual({r.find('a:rPr', NS).get('spc') for r in runs}, {'100'})

    def test_bullets(self):
        paras = self.sh['list'].findall('.//a:p', NS)
        self.assertEqual(len(paras), 2)
        self.assertEqual(len(paras[1].findall('a:br', NS)), 1)  # continuation line in the same paragraph
        ppr = paras[0].find('a:pPr', NS)
        self.assertEqual(ppr.find('a:buChar', NS).get('char'), '●')
        self.assertEqual(ppr.find('a:buClr/a:srgbClr', NS).get('val'), 'F26B3A')
        self.assertEqual(ppr.find('a:buSzPct', NS).get('val'), '125000')
        self.assertEqual(paras[1].find('a:pPr/a:spcBef/a:spcPts', NS).get('val'), '600')

    def test_shapes(self):
        self.assertEqual(self.geom('card').get('prst'), 'roundRect')
        self.assertEqual(self.geom('card').find('a:avLst/a:gd', NS).get('fmla'), 'val 6000')  # 12 / 200
        self.assertEqual(self.geom('pill').find('a:avLst/a:gd', NS).get('fmla'), 'val 50000')
        self.assertEqual(self.sh['pill'].find('.//a:bodyPr', NS).get('anchor'), 'ctr')
        self.assertEqual(self.geom('dot').get('prst'), 'ellipse')
        self.assertEqual(self.sh['dot'].find('.//a:alpha', NS).get('val'), '50000')
        self.assertEqual(self.geom('step1').get('prst'), 'homePlate')
        self.assertEqual(self.geom('step1').find('a:avLst/a:gd', NS).get('fmla'), 'val 33333')  # 20 / 60
        self.assertEqual(self.geom('step2').get('prst'), 'chevron')
        self.assertIsNotNone(self.sh['tri'].find('p:spPr/a:custGeom', NS))
        self.assertEqual(self.sh['card'].find('p:spPr/a:ln', NS).get('w'), str(K))

    def test_line_and_image(self):
        ln = self.sh['rule'].find('p:spPr/a:ln', NS)
        self.assertEqual((ln.get('w'), ln.find('a:prstDash', NS).get('val')), (str(2 * K), 'dash'))
        self.assertIsNotNone(ln.find('a:tailEnd', NS))
        self.assertIsNone(ln.find('a:headEnd', NS))
        self.assertEqual(xfrm_of(self.sh['logo']), (880 * K, 20 * K, 60 * K, 30 * K))
        asset = Image.open(os.path.join(os.path.dirname(self.out), 'assets', 'logo.png'))
        self.assertEqual(asset.size, (180, 90))  # upscaled x3
        alpha = asset.getchannel('A')
        self.assertEqual(alpha.getpixel((2, 2)), 0)        # background made transparent
        self.assertEqual(alpha.getpixel((90, 45)), 255)    # the red mark is kept

    def test_skipped_elements_are_left_out(self):
        self.assertNotIn('ghost', self.sh)

    def test_opens_in_python_pptx(self):
        try:
            from pptx import Presentation
        except ImportError:
            self.skipTest('python-pptx not installed (optional cross-check)')
        prs = Presentation(self.out)
        self.assertEqual(prs.slides[0].shapes[0].text_frame.text, 'Bonjour & <bienvenue>')


if __name__ == '__main__':
    unittest.main()
