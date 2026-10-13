"""pptx_writer.py: the package it writes is a valid, complete .pptx."""
import io
import os
import tempfile
import unittest
import zipfile

from PIL import Image

from helpers import NS, shapes_by_name, validate_pptx, xfrm_of
from pptx_writer import Deck, body_xml, paragraph_xml


def png_bytes(color=(255, 0, 0, 128), size=(8, 8)):
    buf = io.BytesIO()
    Image.new('RGBA', size, color).save(buf, 'PNG')
    return buf.getvalue()


class WriterTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = os.path.join(cls.tmp.name, 'writer.pptx')
        deck = Deck(12192000, 6858000, title='Test & <deck>')
        s = deck.add_slide()
        s.background = '#F7F5F0'
        base = {'font': 'Arial', 'size_pt': 18, 'color': '#112233'}
        p = paragraph_xml([[('Hello <world> & "friends"\x01', dict(base, bold=True)), (' again', base)],
                           [('second line', dict(base, spacing_pt=1.5))]], base, {'align': 'left', 'line_pt': 24}, 'fr-FR')
        s.shape('title', 914400, 457200, 6000000, 800000, textbox=True, text=body_xml([p]))
        bullet = {'mar_emu': 228600, 'color': '#F26B3A', 'size_pct': 125, 'char': '●'}
        b1 = paragraph_xml([[('point one', base)]], base, {'bullet': bullet, 'line_pt': 24}, 'en-US')
        b2 = paragraph_xml([[('point two', base)]], base, {'bullet': bullet, 'line_pt': 24, 'before_pt': 6}, 'en-US')
        s.shape('list', 914400, 1500000, 6000000, 800000, textbox=True, text=body_xml([b1, b2]), group='g1')
        s.shape('card', 1000000, 2500000, 3000000, 1500000, prst='roundRect', adj=[0.1], fill='#FFFFFF',
                line={'color': '#DDDDDD', 'width_emu': 12700}, group='g1')
        s.shape('chev', 5000000, 2500000, 2000000, 600000, prst='chevron', adj=[0.3], fill='#1F2A44', alpha=0.5)
        s.polygon('tri', [(0, 0), (500000, 0), (250000, 400000)], fill='#00FF00')
        s.connector('rule', 7000000, 4000000, 6000000, 4500000,
                    {'color': '#999999', 'width_emu': 9525, 'dash': 'dash', 'tail': True})
        s.picture('icon', png_bytes(), 100000, 100000, 400000, 400000, descr='red "square"')
        s2 = deck.add_slide()
        s2.picture('icon2', png_bytes((0, 0, 255, 255)), 0, 0, 100, 100)
        deck.save(cls.path)
        cls.slides = validate_pptx(cls.path)
        cls.s1 = cls.slides['ppt/slides/slide1.xml']

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_package_parts(self):
        self.assertEqual(sorted(self.slides), ['ppt/slides/slide1.xml', 'ppt/slides/slide2.xml'])
        with zipfile.ZipFile(self.path) as z:
            names = z.namelist()
            self.assertEqual(names[0], '[Content_Types].xml')
            self.assertIn('ppt/media/image1.png', names)
            self.assertIn('ppt/media/image2.png', names)
            core = z.read('docProps/core.xml').decode('utf-8')
        self.assertIn('Test &amp; &lt;deck&gt;', core)

    def test_text_is_escaped_and_kept(self):
        texts = [t.text for t in self.s1.iter(f'{{{NS["a"]}}}t')]
        self.assertIn('Hello <world> & "friends"', texts)  # control character dropped, the rest verbatim
        self.assertIn('second line', texts)
        runs = list(self.s1.iter(f'{{{NS["a"]}}}rPr'))
        self.assertTrue(all(r.get('lang') in ('fr-FR', 'en-US') for r in runs))
        spaced = [r for r in runs if r.get('spc')]
        self.assertEqual([r.get('spc') for r in spaced], ['150'])

    def test_paragraph_properties(self):
        ppr = [p for p in self.s1.iter(f'{{{NS["a"]}}}pPr')]
        self.assertEqual(ppr[0].find('a:lnSpc/a:spcPts', NS).get('val'), '2400')
        self.assertIsNotNone(ppr[0].find('a:buNone', NS))
        b2 = ppr[2]
        self.assertEqual(b2.get('marL'), '228600')
        self.assertEqual(b2.get('indent'), '-228600')
        self.assertEqual(b2.find('a:buChar', NS).get('char'), '●')
        self.assertEqual(b2.find('a:buSzPct', NS).get('val'), '125000')
        self.assertEqual(b2.find('a:spcBef/a:spcPts', NS).get('val'), '600')

    def test_text_boxes_do_not_autofit(self):
        for body in self.s1.iter(f'{{{NS["a"]}}}bodyPr'):
            self.assertIsNotNone(body.find('a:noAutofit', NS))
            self.assertEqual([body.get(k) for k in ('lIns', 'tIns', 'rIns', 'bIns')], ['0'] * 4)

    def test_group_bounds_children(self):
        sh = shapes_by_name(self.s1)
        g = sh['g1']
        self.assertEqual(g.tag, f'{{{NS["p"]}}}grpSp')
        # union of 'list' (914400, 1500000, 6000000 x 800000) and 'card' (1000000, 2500000, 3000000 x 1500000)
        self.assertEqual(xfrm_of(g), (914400, 1500000, 6000000, 4000000 - 1500000))
        self.assertEqual(len(g.findall('p:sp', NS)), 2)

    def test_shapes(self):
        sh = shapes_by_name(self.s1)
        card = sh['card'].find('p:spPr/a:prstGeom', NS)
        self.assertEqual(card.get('prst'), 'roundRect')
        self.assertEqual(card.find('a:avLst/a:gd', NS).attrib, {'name': 'adj', 'fmla': 'val 10000'})
        chev = sh['chev'].find('p:spPr', NS)
        self.assertEqual(chev.find('a:solidFill/a:srgbClr/a:alpha', NS).get('val'), '50000')
        self.assertIsNotNone(sh['tri'].find('p:spPr/a:custGeom/a:pathLst/a:path/a:close', NS))
        rule = sh['rule'].find('p:spPr/a:xfrm', NS)
        self.assertEqual(rule.get('flipH'), '1')
        self.assertIsNone(rule.get('flipV'))
        ln = sh['rule'].find('p:spPr/a:ln', NS)
        self.assertEqual(ln.find('a:prstDash', NS).get('val'), 'dash')
        self.assertIsNotNone(ln.find('a:tailEnd', NS))
        self.assertEqual(sh['icon'].find('p:nvPicPr/p:cNvPr', NS).get('descr'), 'red "square"')

    def test_font_slots_language_and_direction(self):
        from pptx_writer import paragraph_xml, run_xml
        st = {'font': 'Arial', 'font_ea': 'Yu Gothic', 'font_cs': 'Dubai', 'lang': 'ar-SA', 'size_pt': 12}
        r = run_xml('مرحبا', st, 'en-US')
        self.assertIn('lang="ar-SA"', r)  # the run's own language wins over the slide's
        self.assertIn('<a:latin typeface="Arial"/><a:ea typeface="Yu Gothic"/><a:cs typeface="Dubai"/>', r)
        self.assertIn('<a:ea typeface="Arial"/>', run_xml('x', {'font': 'Arial', 'size_pt': 12}, 'en-US'))
        p = paragraph_xml([[('مرحبا', st)]], st, {'align': 'right', 'rtl': True}, 'en-US')
        self.assertIn('<a:pPr algn="r" rtl="1">', p)
        self.assertNotIn('rtl=', paragraph_xml([[('x', st)]], st, {}, 'en-US'))

    def test_flipped_shape(self):
        deck = Deck(12192000, 6858000)
        s = deck.add_slide()
        s.shape('left-arrow', 0, 0, 100, 50, prst='chevron', adj=[0.3], fill='#000000', flip_h=True)
        self.assertIn('<a:xfrm flipH="1">', s.xml())

    def test_background(self):
        bg = self.s1.find('p:cSld/p:bg/p:bgPr/a:solidFill/a:srgbClr', NS)
        self.assertEqual(bg.get('val'), 'F7F5F0')

    def test_opens_in_python_pptx(self):
        try:
            from pptx import Presentation
        except ImportError:
            self.skipTest('python-pptx not installed (optional cross-check)')
        prs = Presentation(self.path)
        self.assertEqual(len(prs.slides), 2)
        self.assertEqual(prs.slide_width, 12192000)
        names = [sh.name for sh in prs.slides[0].shapes]
        self.assertEqual(names, ['title', 'g1', 'chev', 'tri', 'rule', 'icon'])
        self.assertEqual(prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text, 'Hello <world> & "friends"')


if __name__ == '__main__':
    unittest.main()
